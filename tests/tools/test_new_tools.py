"""v1.5.0 新增内置工具的功能与安全测试。

安全用例集中在四个"不能出事"的地方：
  · http_request —— SSRF（私网/回环/元数据/重定向绕过）
  · archive      —— zip-slip 路径穿越
  · db_query     —— 只读约束与表名注入
  · email/im     —— 外发动作的权限档位与群发上限
"""

from __future__ import annotations

import asyncio
import os
import sqlite3
import zipfile
from pathlib import Path

import pytest

from automind.core.types import PermissionTier
from automind.tools._toolkit import BlockedTarget, check_url, safe_extract_path
from automind.tools.collab_tools import CalendarTool, ImIntegrationTool, NotifyTool
from automind.tools.data_tools import ArchiveTool, DbQueryTool, FileSearchTool
from automind.tools.net_tools import HttpRequestTool, WebSearchTool
from automind.tools.office import EmailTool, ExcelTool, PdfTool, WordTool


def run(coro):
    return asyncio.run(coro)


# ── SSRF ────────────────────────────────────────────────────

class TestSsrfGuard:
    @pytest.mark.parametrize("url", [
        "http://127.0.0.1:8765/api/status",     # 本机的 AutoMind 自己
        "http://localhost/admin",
        "http://169.254.169.254/latest/meta-data/",   # 云实例元数据
        "http://metadata.google.internal/",
        "http://10.0.0.5/internal",
        "http://192.168.1.1/",
        "http://[::1]:8000/",
    ])
    def test_private_and_metadata_blocked(self, url):
        with pytest.raises(BlockedTarget):
            check_url(url)

    @pytest.mark.parametrize("url", [
        "file:///etc/passwd", "ftp://example.com/x", "gopher://x/", "data:text/html,x",
    ])
    def test_dangerous_schemes_blocked(self, url):
        with pytest.raises(BlockedTarget):
            check_url(url)

    def test_allow_private_is_opt_in(self):
        # 显式声明要访问内网时才放行
        assert check_url("http://127.0.0.1:9/", allow_private=True)

    def test_tool_reports_block_without_network(self):
        r = run(HttpRequestTool().execute(url="http://169.254.169.254/"))
        assert not r.success
        assert r.output.get("blocked") is True


# ── zip-slip ────────────────────────────────────────────────

class TestArchiveSafety:
    @pytest.mark.parametrize("member", [
        "../evil.txt", "../../etc/passwd", "a/../../../../tmp/x",
    ])
    def test_zip_slip_rejected(self, member, tmp_path):
        with pytest.raises(BlockedTarget):
            safe_extract_path(tmp_path, member)

    def test_normal_member_allowed(self, tmp_path):
        assert safe_extract_path(tmp_path, "sub/ok.txt")

    def test_extract_refuses_traversing_archive(self, tmp_path):
        eviljar = tmp_path / "evil.zip"
        with zipfile.ZipFile(eviljar, "w") as z:
            z.writestr("../escaped.txt", "pwned")
        dest = tmp_path / "out"
        r = run(ArchiveTool().execute(action="extract", path=str(eviljar), dest=str(dest)))
        assert not r.success
        assert not (tmp_path / "escaped.txt").exists(), "文件逃出了解压目录"

    def test_roundtrip(self, tmp_path):
        src = tmp_path / "data.txt"
        src.write_text("hello", encoding="utf-8")
        arc = tmp_path / "a.zip"
        r = run(ArchiveTool().execute(action="create", path=str(arc), sources=[str(src)]))
        assert r.success, r.error
        out = tmp_path / "unpacked"
        r2 = run(ArchiveTool().execute(action="extract", path=str(arc), dest=str(out)))
        assert r2.success, r2.error
        assert (out / "data.txt").read_text(encoding="utf-8") == "hello"


# ── db_query ────────────────────────────────────────────────

@pytest.fixture()
def sample_db(tmp_path):
    db = tmp_path / "t.db"
    con = sqlite3.connect(db)
    con.execute("CREATE TABLE users (id INTEGER, name TEXT)")
    con.executemany("INSERT INTO users VALUES (?,?)", [(1, "amy"), (2, "bob")])
    con.commit()
    con.close()
    return db


class TestDbQuery:
    def test_select_works(self, sample_db):
        r = run(DbQueryTool().execute(database=str(sample_db),
                                      sql="SELECT name FROM users ORDER BY id"))
        assert r.success, r.error
        assert r.output["rows"] == [["amy"], ["bob"]]

    def test_parameter_binding(self, sample_db):
        r = run(DbQueryTool().execute(database=str(sample_db),
                                      sql="SELECT name FROM users WHERE id = ?",
                                      params=[2]))
        assert r.success, r.error
        assert r.output["rows"] == [["bob"]]

    def test_write_rejected_in_community(self, sample_db):
        r = run(DbQueryTool().execute(database=str(sample_db),
                                      sql="DELETE FROM users"))
        assert not r.success
        # 社区版：要么被 data_pro 门控挡下，要么被只读连接挡下 —— 都不能删成
        con = sqlite3.connect(sample_db)
        assert con.execute("SELECT COUNT(*) FROM users").fetchone()[0] == 2
        con.close()

    def test_table_name_injection_rejected(self, sample_db):
        r = run(DbQueryTool().execute(database=str(sample_db), action="schema",
                                      table="users); DROP TABLE users;--"))
        assert not r.success
        assert "非法" in (r.error or "")

    def test_tables_listing(self, sample_db):
        r = run(DbQueryTool().execute(database=str(sample_db), action="tables"))
        assert r.success, r.error
        assert any(t["name"] == "users" for t in r.output["tables"])


# ── file_search ─────────────────────────────────────────────

class TestFileSearch:
    def test_finds_by_glob_and_content(self, tmp_path):
        (tmp_path / "a.py").write_text("def hello():\n    return 1\n", encoding="utf-8")
        (tmp_path / "b.txt").write_text("hello world", encoding="utf-8")
        t = FileSearchTool(project_root=tmp_path)
        r = run(t.execute(pattern="*.py"))
        assert r.success and r.output["count"] == 1
        r2 = run(t.execute(pattern="*.py", contains=r"def \w+"))
        assert r2.success and r2.output["matches"][0]["line"] == 1

    def test_path_escape_rejected(self, tmp_path):
        r = run(FileSearchTool(project_root=tmp_path).execute(path="../.."))
        assert not r.success
        assert "超出" in (r.error or "")


# ── 办公工具（无第三方库时应给出可操作提示，而不是崩） ──────

class TestOfficeTools:
    def test_excel_roundtrip_or_clear_hint(self, tmp_path):
        p = tmp_path / "b.xlsx"
        r = run(ExcelTool().execute(action="create", path=str(p),
                                    rows=[["名称", "数量"], ["苹果", 3]]))
        if not r.success:
            assert "pip install" in (r.error or ""), r.error
            pytest.skip("openpyxl 未安装 —— 已验证给出安装提示")
        r2 = run(ExcelTool().execute(action="read", path=str(p)))
        assert r2.success, r2.error
        assert r2.output["rows"][1] == ["苹果", 3]

    def test_word_roundtrip_or_clear_hint(self, tmp_path):
        p = tmp_path / "d.docx"
        r = run(WordTool().execute(action="create", path=str(p),
                                   heading="报告", paragraphs=["第一段"]))
        if not r.success:
            assert "pip install" in (r.error or ""), r.error
            pytest.skip("python-docx 未安装 —— 已验证给出安装提示")
        r2 = run(WordTool().execute(action="read", path=str(p)))
        assert r2.success and "第一段" in r2.output["paragraphs"]

    def test_word_rejects_legacy_doc(self, tmp_path):
        r = run(WordTool().execute(action="read", path=str(tmp_path / "x.doc")))
        assert not r.success
        # 缺库时先报缺库也可接受，但不能假装支持 .doc
        assert (".docx" in (r.error or "")) or ("pip install" in (r.error or ""))

    def test_pdf_missing_file(self, tmp_path):
        r = run(PdfTool().execute(action="info", path=str(tmp_path / "none.pdf")))
        assert not r.success

    def test_pdf_page_range_parsing(self):
        assert PdfTool._parse_pages("1-3", 10) == [0, 1, 2]
        assert PdfTool._parse_pages("2,5", 10) == [1, 4]
        # 3 页的文档要第 9-99 页 —— 一页都不该返回（而不是硬凑一页出来）
        assert PdfTool._parse_pages("9-99", 3) == []
        assert PdfTool._parse_pages("2-99", 3) == [1, 2]   # 部分越界只截断超出部分


# ── 外发动作的权限与上限 ────────────────────────────────────

class TestOutboundGuards:
    def test_email_send_is_dangerous_tier(self):
        # DANGEROUS ⇒「询问」模式下逐封审批，且审批失败一律拒绝（v1.4.5）
        assert EmailTool().permission_tier is PermissionTier.DANGEROUS

    def test_im_send_is_dangerous_tier(self):
        assert ImIntegrationTool().permission_tier is PermissionTier.DANGEROUS

    def test_mass_mailing_refused(self, monkeypatch):
        monkeypatch.setenv("AUTOMIND_SMTP_HOST", "smtp.example.com")
        monkeypatch.setenv("AUTOMIND_SMTP_USER", "me@example.com")
        monkeypatch.setenv("AUTOMIND_SMTP_PASSWORD", "x")
        many = [f"u{i}@example.com" for i in range(50)]
        r = run(EmailTool().execute(action="send", to=many, subject="s", body="b"))
        assert not r.success
        assert "群发" in (r.error or "")

    def test_email_password_never_taken_from_args(self, monkeypatch):
        monkeypatch.delenv("AUTOMIND_SMTP_HOST", raising=False)
        r = run(EmailTool().execute(action="send", to=["a@b.com"], subject="s",
                                    body="b", password="hunter2"))
        assert not r.success
        assert "AUTOMIND_SMTP_HOST" in (r.error or "")

    def test_im_requires_webhook_config(self, monkeypatch):
        monkeypatch.delenv("AUTOMIND_IM_WEBHOOK", raising=False)
        r = run(ImIntegrationTool().execute(text="hi"))
        assert not r.success and r.output.get("needs_config")

    def test_im_webhook_ssrf_checked(self, monkeypatch):
        monkeypatch.setenv("AUTOMIND_IM_WEBHOOK", "http://127.0.0.1:8765/api/run")
        r = run(ImIntegrationTool().execute(text="hi"))
        assert not r.success and r.output.get("blocked")

    def test_websearch_needs_config(self, monkeypatch):
        monkeypatch.delenv("AUTOMIND_SEARCH_PROVIDER", raising=False)
        r = run(WebSearchTool().execute(query="x"))
        assert not r.success and r.output.get("needs_config")


# ── 日历 ────────────────────────────────────────────────────

class TestCalendar:
    def test_ics_roundtrip_or_clear_hint(self, tmp_path):
        p = tmp_path / "c.ics"
        r = run(CalendarTool().execute(action="create", path=str(p), summary="周会",
                                       start="2026-08-10T14:00:00", duration_minutes=30))
        if not r.success:
            assert "pip install" in (r.error or ""), r.error
            pytest.skip("icalendar 未安装 —— 已验证给出安装提示")
        r2 = run(CalendarTool().execute(action="list", path=str(p), days_ahead=3650))
        assert r2.success, r2.error
        assert any(e["summary"] == "周会" for e in r2.output["events"])

    def test_outlook_rejected_off_windows(self):
        import platform
        if platform.system() == "Windows":
            pytest.skip("本用例断言的是非 Windows 平台的行为")
        r = run(CalendarTool().execute(action="outlook_list"))
        assert not r.success and "Windows" in (r.error or "")


# ── 通知（不实际弹窗，只验证参数转义正确） ──────────────────

class TestNotify:
    def test_powershell_quote_escaping(self):
        # 标题里带单引号不能破坏 PowerShell 字面量 —— 否则就是命令注入
        assert NotifyTool._ps_str("it's") == "'it''s'"
        # 别忘了外层还要再包一对引号。已实测该字面量在 PowerShell 里求值为
        # 纯数据（VALUE=['; Write-Output INJECTED; ']），注入不会被执行。
        assert NotifyTool._ps_str("'; rm -rf /; '") == "'''; rm -rf /; '''"

    def test_tool_is_safe_tier(self):
        assert NotifyTool().permission_tier is PermissionTier.SAFE


# ── 注册与门控 ──────────────────────────────────────────────

class TestRegistrationAndGating:
    def test_all_new_tools_registered(self):
        from automind.agent import AutoMindAgent
        from automind.core.config import AgentConfig
        a = AutoMindAgent(AgentConfig())
        names = {t.name for t in a.tool_registry._tools.values()}
        expected = {"excel_tool", "word_tool", "pdf_tool", "email_tool",
                    "web_search", "http_request", "db_query", "file_search",
                    "archive", "notify", "calendar", "im_integration"}
        assert expected <= names, f"未注册：{expected - names}"

    def test_pro_actions_gated_in_community(self, tmp_path):
        """社区版下进阶动作必须被挡，并给出升级引导。"""
        p = tmp_path / "b.xlsx"
        r = run(ExcelTool().execute(action="style", path=str(p)))
        assert not r.success
        assert r.output.get("upgrade_required") == "office_pro"

    def test_basic_actions_not_gated(self, tmp_path):
        """基础动作不能被误伤 —— 社区版要真的能用。"""
        r = run(ExcelTool().execute(action="read", path=str(tmp_path / "none.xlsx")))
        # 文件不存在是预期的失败，但不应是"需要升级"
        assert r.output.get("upgrade_required") is None

    def test_feature_keys_declared(self):
        from automind.core.edition import COMMERCIAL_FEATURES
        for key in ("office_pro", "data_pro", "integration_pro"):
            assert key in COMMERCIAL_FEATURES


def test_env_isolation_sanity():
    """确保测试没有把真实凭据带进来（本地开发机上跑时的护栏）。"""
    assert not os.environ.get("AUTOMIND_SMTP_PASSWORD_REAL")
    assert Path.cwd().exists()


def _office_pro_or_skip():
    """取 pro 侧的 OfficeProFeature；取不到就跳过。

    ``pro/`` 是闭源商业包且已被 .gitignore 排除 —— 社区版仓库的克隆体与 CI 上
    根本没有这个目录。这些用例只在**同时拥有社区核心与商业包**的开发机上有意义，
    因此按"缺则跳过"处理，而不是让社区版 CI 红一片。
    """
    import sys
    pro_dir = Path(__file__).resolve().parents[2] / "pro"
    if not (pro_dir / "automind_pro" / "office_pro.py").is_file():
        pytest.skip("未随附商业包 automind_pro（社区版仓库的正常状态）")
    if str(pro_dir) not in sys.path:
        sys.path.insert(0, str(pro_dir))
    return pytest.importorskip("automind_pro.office_pro").OfficeProFeature()


class TestProDelegation:
    """进阶动作必须真的能跑通 —— 不能是"授权了却提示不支持"的空壳。"""

    @pytest.fixture()
    def with_office_pro(self, monkeypatch):
        """把 pro 的 OfficeProFeature 注册进来，模拟已授权的专业版。"""
        from automind.core import edition
        monkeypatch.setitem(edition._state["features"], "office_pro", _office_pro_or_skip())
        monkeypatch.setitem(edition._state, "loaded", True)
        yield

    def test_excel_style_runs_for_real(self, tmp_path, with_office_pro):
        p = tmp_path / "s.xlsx"
        assert run(ExcelTool().execute(action="create", path=str(p),
                                       rows=[["名称", "数量"], ["苹果", 3]])).success
        r = run(ExcelTool().execute(action="style", path=str(p), cell_range="A1:B1",
                                    bold=True, fill="FFFF00", autofit=True,
                                    freeze_header=True))
        assert r.success, r.error
        assert r.output["styled_cells"] == 2
        # 落盘确实生效
        import openpyxl
        ws = openpyxl.load_workbook(p).active
        assert ws["A1"].font.bold and ws.freeze_panes == "A2"

    def test_word_template_mail_merge(self, tmp_path, with_office_pro):
        tpl = tmp_path / "t.docx"
        assert run(WordTool().execute(action="create", path=str(tpl),
                                      paragraphs=["尊敬的 {{name}}，您的订单 {{no}} 已发出。"])).success
        out = tmp_path / "filled.docx"
        r = run(WordTool().execute(action="template", path=str(tpl), output=str(out),
                                   values={"name": "张三", "no": "A-1001"}))
        assert r.success, r.error
        assert r.output["replaced"] == 2
        import docx
        text = "\n".join(p.text for p in docx.Document(str(out)).paragraphs)
        assert "张三" in text and "A-1001" in text and "{{" not in text

    def test_pdf_encrypt_then_needs_password(self, tmp_path, with_office_pro):
        import pypdf
        src = tmp_path / "a.pdf"
        w = pypdf.PdfWriter()
        w.add_blank_page(width=200, height=200)
        with open(src, "wb") as f:
            w.write(f)
        r = run(PdfTool().execute(action="encrypt", path=str(src), password="s3cret!"))
        assert r.success, r.error
        assert pypdf.PdfReader(r.output["path"]).is_encrypted
        assert "s3cret!" not in str(r.output), "返回值里不能回显密码"

    def test_community_still_blocked(self, tmp_path):
        """没注册 office_pro 时仍应是升级引导，而不是崩溃或静默成功。"""
        r = run(ExcelTool().execute(action="style", path=str(tmp_path / "x.xlsx")))
        assert not r.success
        assert r.output.get("upgrade_required") == "office_pro"

    def test_every_declared_pro_action_is_implemented(self):
        """声明为进阶的动作，pro 侧必须都有对应实现 —— 防止再出现占位卖点。"""
        from automind.tools.office import excel_tool, pdf_tool, word_tool
        f = _office_pro_or_skip()
        for mod, tool in ((excel_tool, "excel_tool"), (word_tool, "word_tool"),
                          (pdf_tool, "pdf_tool")):
            for action in mod.PRO_ACTIONS:
                assert f.supports(tool, action), f"{tool}.{action} 被标为进阶却无实现"
                assert hasattr(f, f"_{tool}_{action}"), f"{tool}.{action} 缺少处理函数"
