"""artifact_selfcheck —— 写成功 ≠ 落盘正确（内置插件）。

**为什么需要它**：工具的"成功"只说明它没抛异常。
- ``file_write`` 返回的 ``size`` 是**内存里的长度**，不回读磁盘；
- Excel/Word/PPT 保存后**不重开自检**，文件被占用写坏、内容被截断都可能算成功；
- 执行器侧的产物断言只查 ``exists()`` + 大小 > 1 字节，而且只在特定模式下触发。

**做法**：分派中间件在写类工具返回成功后立刻回读磁盘：存在、非空、
Office/PDF/图片能重新打开（缺对应库时**跳过**该项检查，绝不因缺库误判失败）。
不合格就返回 ``success=False`` 的 ToolResult，让模型在下一步修。

**边界**：只覆盖"能定位到文件路径"的写类工具；不做业务语义校验，也不跑测试。
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from automind.core.artifacts import resolve_artifact_path
from automind.core.hooks import AgentHooks
from automind.core.logging import get_logger
from automind.core.types import ToolResult

logger = get_logger("automind.plugin.artifact_selfcheck")

#: 会写盘、且返回值/参数里能定位到路径的工具
_WRITE_TOOLS = frozenset({
    "file_write", "file_edit", "file_multi_edit",
    "excel_tool", "word_tool", "ppt_tool", "pdf_tool", "csv_tool",
    "image_tool", "chart_tool",
})

#: 超过这个体积不做"能否打开"的深度检查（只查存在与非空），避免拖慢任务
_MAX_OPEN_BYTES = 64 * 1024 * 1024


class _MissingLib(Exception):
    """用于区分"缺依赖库"与"文件真的打不开"。"""


def _open_xlsx(p: Path) -> None:
    try:
        import openpyxl
    except ImportError as e:
        raise _MissingLib("openpyxl") from e
    openpyxl.load_workbook(str(p), read_only=True, data_only=True).close()


def _open_docx(p: Path) -> None:
    try:
        import docx
    except ImportError as e:
        raise _MissingLib("python-docx") from e
    docx.Document(str(p))


def _open_pptx(p: Path) -> None:
    try:
        import pptx
    except ImportError as e:
        raise _MissingLib("python-pptx") from e
    pptx.Presentation(str(p))


def _open_pdf(p: Path) -> None:
    try:
        import pypdf
    except ImportError as e:
        raise _MissingLib("pypdf") from e
    reader = pypdf.PdfReader(str(p))
    if len(reader.pages) == 0:
        raise ValueError("PDF 里没有任何页面")


def _open_image(p: Path) -> None:
    try:
        from PIL import Image
    except ImportError as e:
        raise _MissingLib("pillow") from e
    with Image.open(str(p)) as img:
        img.verify()


#: 后缀 → 打开方式（只列"有明确解析器"的格式）
_OPENERS = {
    ".xlsx": _open_xlsx, ".xlsm": _open_xlsx,
    ".docx": _open_docx,
    ".pptx": _open_pptx,
    ".pdf": _open_pdf,
    ".png": _open_image, ".jpg": _open_image, ".jpeg": _open_image,
    ".webp": _open_image, ".bmp": _open_image, ".gif": _open_image,
}


def _problems(p: Path) -> list[str]:
    """回读磁盘，返回发现的问题（空列表 = 通过）。"""
    try:
        if not p.exists():
            return ["文件不存在"]
        if not p.is_file():
            return ["不是一个文件"]
        st = p.stat()
        if st.st_size <= 0:
            return ["文件是 0 字节"]
    except OSError as e:
        return [f"无法读取文件属性：{type(e).__name__}: {e}"]

    opener = _OPENERS.get(p.suffix.lower())
    if opener is None or st.st_size > _MAX_OPEN_BYTES:
        return []                                   # 没解析器 / 太大：只查存在与非空
    try:
        opener(p)
    except _MissingLib as e:
        logger.debug("artifact_selfcheck skipped: missing lib", module=str(e))
        return []                                   # 缺库不能算文件坏了
    except Exception as e:
        return [f"文件打不开：{type(e).__name__}: {e}"]
    return []


def get_middleware() -> Any:
    """写类工具成功后立刻回读校验；不合格就把这次调用改判失败。"""

    async def mw(tool_name: str, kwargs: dict[str, Any], call_next: Any) -> Any:
        result = await call_next(kwargs)
        try:
            if tool_name not in _WRITE_TOOLS or not getattr(result, "success", False):
                return result
            # v1.11.13：用"工具返回值里的绝对路径 + 会话副本/项目根"解析，
            # 而不是拿模型给的相对路径按服务进程 cwd 去查 —— 目录级隔离下
            # 旧写法会把每一次成功写入都误判成"文件不存在"。
            path = resolve_artifact_path(kwargs, result)
            if path is None:
                return result
            problems = _problems(path)
            if not problems:
                return result
            path_text = str(path)
            logger.warning("artifact_selfcheck_failed", tool=tool_name,
                           path=path_text, problems=problems)
            return ToolResult(
                tool_name=getattr(result, "tool_name", "") or tool_name,
                success=False,
                output=getattr(result, "output", None),
                error=(f"写入后自检未通过：{path_text} —— " + "；".join(problems)
                       + "。请检查路径是否被占用/磁盘是否已满，修正后重新生成；"
                         "不要原样重试同一次调用。"),
                metadata={"artifact_selfcheck": True, "problems": problems,
                          "path": path_text},
            )
        except Exception as e:                        # 自检故障不影响原结果
            logger.warning("artifact_selfcheck_error", tool=tool_name,
                           error=f"{type(e).__name__}: {e}")
            return result

    return mw


def get_hooks() -> AgentHooks:
    return AgentHooks()
